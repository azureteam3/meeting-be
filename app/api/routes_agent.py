"""
app/api/routes_agent.py - 에이전트 API 라우트 (명세서 §4).

POST /agent/sessions/{id}/query            채팅 질의 (요약·검색·Q&A)
GET  /agent/sessions/{id}/summary          현재 요약 (없으면 생성·저장)
GET  /agent/sessions/{id}/action-items     액션 아이템
GET  /agent/sessions/{id}/decisions        의사결정
GET  /agent/sessions/{id}/minutes          회의록 (없으면 생성·저장)
GET  /agent/sessions/{id}/minutes/download 회의록 .docx 다운로드

※ session_id(경로)는 회의를 식별하는 값 = 자막(transcript_records)의 meeting_id 기준으로 조회.
"""

import io
import uuid
from datetime import date
from urllib.parse import quote

from docx import Document
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.db.postgres import SessionLocal
from app.repositories import agent_repository as repository
from app.schemas.agent import QueryRequest, QueryResponse
from app.services import agent_service as agent


def get_db():
    """요청마다 DB 세션을 빌려주고 끝나면 닫는다.
    (dev3 postgres.py엔 get_db가 없어서, 공용 SessionLocal로 여기서 자체 정의)"""
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


router = APIRouter(prefix="/agent", tags=["agent"])


def _classify_intent(message: str) -> str:
    """질문 종류를 대략 분류 (명세서 intent 값 · 감지 키워드는 명세서 4.2 표 기준)."""
    lower = message.lower()
    if "요약" in message or "summar" in lower or "まとめ" in message or "要約" in message or "总结" in message:
        return "summary"
    if "액션" in message or "할 일" in message or "action" in lower:
        return "action_items"
    if "찾" in message or "검색" in message:
        return "search"
    return "qa"


# 요약 JSON을 채팅 말풍선 텍스트로 바꿀 때 쓰는 언어별 제목들
_SUMMARY_LABELS = {
    "ko": ("📋 회의 요약", "핵심 논의", "결정사항", "액션 아이템", "미결 주제"),
    "en": ("📋 Meeting Summary", "Key Discussions", "Decisions", "Action Items", "Pending Topics"),
    "ja": ("📋 会議要約", "主な議論", "決定事項", "アクションアイテム", "未決トピック"),
    "zh": ("📋 会议摘要", "主要讨论", "决定事项", "行动项目", "未决主题"),
}


def _render_summary_answer(summary_json: dict, lang: str = "ko") -> str:
    """저장된 요약 JSON → 채팅 답변 텍스트. (gpt 추가 호출 없이 파이썬으로 변환)"""
    title, kd, dec, ai, pt = _SUMMARY_LABELS.get(lang.split("-")[0], _SUMMARY_LABELS["ko"])
    lines = [title]
    if summary_json.get("keyDiscussions"):
        lines.append(f"\n{kd}")
        lines += [f"• {x}" for x in summary_json["keyDiscussions"]]
    if summary_json.get("decisions"):
        lines.append(f"\n{dec}")
        for d in summary_json["decisions"]:
            item = f"• {d.get('item', '')}: {d.get('decision', '')}"
            if d.get("decidedBy"):
                item += f" ({d['decidedBy']})"
            lines.append(item)
    if summary_json.get("actionItems"):
        lines.append(f"\n{ai}")
        lines += [f"• {a.get('assignee', '')}: {a.get('task', '')}" for a in summary_json["actionItems"]]
    if summary_json.get("pendingTopics"):
        lines.append(f"\n{pt}")
        lines += [f"• {x}" for x in summary_json["pendingTopics"]]
    if len(lines) == 1 and summary_json.get("overview"):   # 비정형(폴백) 요약이면 원문이라도
        lines.append(str(summary_json["overview"]))
    return "\n".join(lines)


@router.post("/sessions/{session_id}/query", response_model=QueryResponse)
def query(session_id: str, body: QueryRequest, db: Session = Depends(get_db)):
    """채팅 질의(요약·검색·Q&A)를 처리한다. (명세서 4.2)

    - intent가 summary(누적 전체 요약)면: 스냅샷 파이프라인을 태워 증분 갱신 + DB 저장까지 하고
      그 요약을 채팅 답변으로 변환해 준다 (= 설계의 REALTIME_SUMMARY. gpt 추가 호출 없어 빠름).
    - 그 외(검색·Q&A 등)는 멀티툴 에이전트가 답한다 (= CHAT_ANSWER. 저장 안 함).
    """
    intent = _classify_intent(body.message)

    if intent == "summary":
        snap = _get_fresh_summary(db, session_id, body.responseLanguage)
        if snap is not None:
            return QueryResponse(
                queryId="qry-" + uuid.uuid4().hex[:12],
                intent=intent,
                answer=_render_summary_answer(snap.summary_json, body.responseLanguage),
            )
        # 자막이 아예 없으면 아래 일반 채팅으로 넘어가 에이전트가 상황을 안내한다

    answer = agent.answer_question(db, session_id, body.message, body.responseLanguage)
    return QueryResponse(
        queryId="qry-" + uuid.uuid4().hex[:12],
        intent=intent,
        answer=answer,
    )


def _get_fresh_summary(db: Session, session_id: str, lang: str = "ko", refresh: bool = False):
    """요약 스냅샷을 '신선한 상태'로 가져온다. (summary·action-items·decisions 공용)

    - refresh=True          : 처음부터 전체 재요약
    - 저장본 없음            : 전체 요약 생성·저장
    - 새 자막 도착(워터마크)  : 기존 요약 + 새 자막만으로 증분 갱신·저장
    - 그 외                 : 저장본 그대로 (gpt 호출 없음)
    자막이 아예 없으면 None.
    """
    snap = None if refresh else repository.get_summary_snapshot(db, session_id)
    last_id = repository.get_last_transcript_id(db, session_id)
    if snap is None or (last_id is not None and last_id > snap.coverage_to_seq):
        snap, _count = agent.generate_and_save_summary(db, session_id, lang, prev_snap=snap)
    return snap


@router.get("/sessions/{session_id}/summary")
def get_summary(
    session_id: str,
    lang: str = "ko",
    refresh: bool = False,     # true면 저장본 무시하고 지금 자막으로 새로 만든다
    db: Session = Depends(get_db),
):
    """현재 시점 요약. (명세서 4.3)

    - 저장본이 없으면: 전체 자막으로 생성·저장
    - 새 자막이 도착했으면(워터마크 비교): 기존 요약 + 새 자막만으로 '증분 갱신'
    - 그 외: 저장본 그대로 (gpt 호출 없음)
    - refresh=true: 처음부터 전체 재요약 (증분 드리프트가 의심될 때 수동 초기화)
    """
    snap = _get_fresh_summary(db, session_id, lang, refresh)
    if snap is None:
        raise HTTPException(status_code=404, detail="no_transcript")

    return {
        "sessionId": str(session_id),
        "generatedAt": snap.created_at,
        "version": snap.version,
        "coverageToSeq": snap.coverage_to_seq,
        **snap.summary_json,   # keyDiscussions / decisions / actionItems / pendingTopics
    }


@router.get("/sessions/{session_id}/action-items")
def get_action_items(session_id: str, db: Session = Depends(get_db)):
    """액션 아이템만 따로 반환. (명세서 4.4) 요약이 낡았으면 증분 갱신 후 꺼낸다."""
    snap = _get_fresh_summary(db, session_id)
    if snap is None:
        raise HTTPException(status_code=404, detail="no_transcript")
    items = snap.summary_json.get("actionItems", [])
    return {"total": len(items), "items": items}


@router.get("/sessions/{session_id}/decisions")
def get_decisions(session_id: str, db: Session = Depends(get_db)):
    """의사결정만 따로 반환. (명세서 4.5) 요약이 낡았으면 증분 갱신 후 꺼낸다."""
    snap = _get_fresh_summary(db, session_id)
    if snap is None:
        raise HTTPException(status_code=404, detail="no_transcript")
    decisions = snap.summary_json.get("decisions", [])
    return {"total": len(decisions), "decisions": decisions}


def _build_minutes_docx(data: dict) -> io.BytesIO:
    """회의록 dict를 Word(.docx) 파일(메모리 버퍼)로 만든다."""
    doc = Document()
    doc.add_heading(data.get("title") or "회의록", level=0)
    doc.add_paragraph(f"날짜: {date.today().isoformat()}")
    doc.add_paragraph(f"참석자: {', '.join(data.get('participants') or []) or '-'}")

    doc.add_heading("핵심 논의", level=1)
    for d in data.get("keyDiscussions") or []:
        doc.add_paragraph(str(d), style="List Bullet")

    doc.add_heading("결정 사항", level=1)
    for x in data.get("decisions") or []:
        line = f"{x.get('item', '')}: {x.get('decision', '')}"
        if x.get("decidedBy"):
            line += f" (결정: {x['decidedBy']})"
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("액션 아이템", level=1)
    for a in data.get("actionItems") or []:
        line = f"{a.get('assignee', '')}: {a.get('task', '')}"
        if a.get("due"):
            line += f" (기한: {a['due']})"
        doc.add_paragraph(line, style="List Bullet")

    pending = data.get("pendingTopics") or []
    if pending:
        doc.add_heading("미결 주제", level=1)
        for p in pending:
            doc.add_paragraph(str(p), style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


@router.get("/sessions/{session_id}/minutes/download")
def download_minutes(session_id: str, lang: str = "ko", db: Session = Depends(get_db)):
    """전체 자막 + 저장된 요약을 gpt로 회의록으로 정리해 .docx 파일로 다운로드시킨다."""
    data = agent.generate_minutes(db, session_id, lang)
    if data is None:
        raise HTTPException(status_code=404, detail="no_transcript")

    repository.save_minutes(db, session_id, data)   # 다운로드한 회의록도 DB에 이력으로 남긴다

    buf = _build_minutes_docx(data)
    kor_name = f"회의록_{date.today().isoformat()}.docx"
    disposition = (
        f"attachment; filename=meeting_minutes.docx; "
        f"filename*=UTF-8''{quote(kor_name)}"
    )
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": disposition},
    )


def _minutes_response(row) -> dict:
    """회의록 DB 행 -> 응답 dict."""
    return {
        "minutesId": str(row.id),
        "sessionId": str(row.meeting_id),
        "createdAt": row.created_at,
        **row.minutes_json,
    }


@router.get("/sessions/{session_id}/minutes")
def get_or_create_minutes(
    session_id: str,
    lang: str = "ko",
    refresh: bool = False,     # true면 저장본 무시하고 새로 만들어 저장
    db: Session = Depends(get_db),
):
    """회의록 반환. 저장본 있으면 그대로, 없으면(또는 refresh) 만들어 저장 후 반환. (명세서 4.6 대체·동기)"""
    row = None if refresh else repository.get_latest_minutes(db, session_id)
    if row is None:
        data = agent.generate_minutes(db, session_id, lang)
        if data is None:
            raise HTTPException(status_code=404, detail="no_transcript")
        row = repository.save_minutes(db, session_id, data)
    return _minutes_response(row)


@router.get("/sessions/{session_id}/minutes/{minutes_id}")
def get_minutes_by_id(session_id: str, minutes_id: uuid.UUID, db: Session = Depends(get_db)):
    """저장된 회의록을 id(minutesId)로 조회한다. (명세서 4.7) 없으면 404."""
    row = repository.get_minutes(db, minutes_id)
    if row is None:
        raise HTTPException(status_code=404, detail="minutes_not_found")
    return _minutes_response(row)
